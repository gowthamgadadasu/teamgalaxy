# backend/attendance_service.py
import csv
import io
import os
import uuid
from datetime import datetime
import openpyxl
from backend.db import get_db

REQUIRED_HEADERS = ["Roll Number", "Classes Conducted", "Classes Attended"]
TEMPLATE_COLUMNS = [
    "Roll Number", "Student Name", "Department", "Year / Section",
    "Classes Conducted", "Classes Attended", "Attendance Percentage", "Data As Of"
]

def get_current_threshold():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT value FROM system_config WHERE key = 'attendance_threshold'")
    row = cursor.fetchone()
    conn.close()
    if row:
        try:
            return float(row["value"])
        except ValueError:
            return 75.0
    return 75.0

def set_threshold(new_threshold, updated_by="In-Charge"):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
    INSERT OR REPLACE INTO system_config (key, value, updated_at, updated_by)
    VALUES ('attendance_threshold', ?, datetime('now', 'localtime'), ?)
    """, (str(float(new_threshold)), updated_by))
    conn.commit()
    conn.close()
    return float(new_threshold)

def generate_csv_template():
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(TEMPLATE_COLUMNS)
    writer.writerow(["24AI001", "Aarav Reddy", "AI", "II / A", "100", "88", "88.0%", "2026-10-10"])
    writer.writerow(["24CS002", "Lavanya Gupta", "CSE", "II / A", "100", "75", "75.0%", "2026-10-10"])
    writer.writerow(["23EC004", "Sneha Pillai", "ECE", "III / A", "96", "66", "68.8%", "2026-10-10"])
    return output.getvalue()

def generate_xlsx_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Template"
    ws.append(TEMPLATE_COLUMNS)
    ws.append(["24AI001", "Aarav Reddy", "AI", "II / A", 100, 88, "88.0%", "2026-10-10"])
    ws.append(["24CS002", "Lavanya Gupta", "CSE", "II / A", 100, 75, "75.0%", "2026-10-10"])
    ws.append(["23EC004", "Sneha Pillai", "ECE", "III / A", 96, 66, "68.8%", "2026-10-10"])
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()

def parse_file_to_rows(file_storage, filename):
    rows = []
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".csv":
        content = file_storage.read().decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(content))
        # normalize headers
        field_map = {}
        if reader.fieldnames:
            for f in reader.fieldnames:
                clean_f = f.strip()
                field_map[clean_f.lower()] = f

            for row_idx, raw_row in enumerate(reader, start=2):
                normalized = {}
                for k, v in raw_row.items():
                    if k:
                        normalized[k.strip()] = v.strip() if isinstance(v, str) else v
                rows.append({"row_number": row_idx, "raw": normalized})

    elif ext in [".xlsx", ".xls"]:
        wb = openpyxl.load_workbook(file_storage, data_only=True)
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        if not all_rows:
            return []
        headers = [str(h).strip() if h is not None else "" for h in all_rows[0]]
        for row_idx, row_values in enumerate(all_rows[1:], start=2):
            if not any(row_values):
                continue
            row_dict = {}
            for col_idx, h in enumerate(headers):
                if h and col_idx < len(row_values):
                    val = row_values[col_idx]
                    row_dict[h] = str(val).strip() if val is not None else ""
            rows.append({"row_number": row_idx, "raw": row_dict})

    return rows

def validate_attendance_preview(rows):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT roll_number, name, dept FROM users WHERE role = 'student'")
    existing_students = {r["roll_number"].upper(): dict(r) for r in cursor.fetchall()}
    conn.close()

    threshold = get_current_threshold()
    now_dt = datetime(2026, 10, 12) # consistent reference date

    valid_records = []
    invalid_records = []
    unmatched_records = []
    duplicates = []
    seen_rolls = {}

    for item in rows:
        row_num = item["row_number"]
        raw = item["raw"]

        # Find roll number key flexibly
        roll_key = next((k for k in raw if "roll" in k.lower()), None)
        conducted_key = next((k for k in raw if "conducted" in k.lower()), None)
        attended_key = next((k for k in raw if "attended" in k.lower()), None)
        name_key = next((k for k in raw if "name" in k.lower()), None)
        dept_key = next((k for k in raw if "dept" in k.lower()), None)
        as_of_key = next((k for k in raw if "as of" in k.lower() or "date" in k.lower()), None)

        errors = []

        roll = str(raw.get(roll_key, "")).strip().upper() if roll_key else ""
        name = str(raw.get(name_key, "")).strip() if name_key else ""
        dept = str(raw.get(dept_key, "")).strip() if dept_key else ""
        as_of = str(raw.get(as_of_key, "2026-10-10")).strip() if as_of_key else "2026-10-10"

        if not roll:
            errors.append("Missing Roll Number")
        
        # Check duplicate in file
        if roll:
            if roll in seen_rolls:
                errors.append(f"Duplicate Roll Number '{roll}' found (first at row {seen_rolls[roll]})")
                duplicates.append({"roll": roll, "row": row_num})
            else:
                seen_rolls[roll] = row_num

        # Check matched student
        is_matched = False
        if roll:
            if roll in existing_students:
                is_matched = True
                if not name:
                    name = existing_students[roll]["name"]
                if not dept:
                    dept = existing_students[roll]["dept"]
            else:
                unmatched_records.append({"roll": roll, "row": row_num})
                errors.append(f"Roll Number '{roll}' not matched to any registered student profile")

        # Validate conducted count
        conducted_raw = raw.get(conducted_key, "") if conducted_key else ""
        attended_raw = raw.get(attended_key, "") if attended_key else ""

        conducted_val = None
        attended_val = None

        try:
            conducted_val = float(str(conducted_raw).replace(",", "").strip())
            if conducted_val < 0:
                errors.append("Classes Conducted cannot be negative")
        except (ValueError, TypeError):
            errors.append(f"Invalid non-numeric Classes Conducted value '{conducted_raw}'")

        try:
            attended_val = float(str(attended_raw).replace(",", "").strip())
            if attended_val < 0:
                errors.append("Classes Attended cannot be negative")
        except (ValueError, TypeError):
            errors.append(f"Invalid non-numeric Classes Attended value '{attended_raw}'")

        if conducted_val is not None and attended_val is not None:
            if attended_val > conducted_val:
                errors.append(f"Classes Attended ({attended_val}) cannot exceed Classes Conducted ({conducted_val})")

        if errors:
            invalid_records.append({
                "row": row_num,
                "rollNumber": roll,
                "name": name,
                "errors": errors,
                "raw": raw
            })
        else:
            # Calculate percentage & eligibility
            if conducted_val == 0:
                pct = None
                elig_status = "Unavailable"
                elig_reason = "Zero classes conducted"
            else:
                pct = round((attended_val / conducted_val) * 100.0, 2)
                
                # Check staleness
                is_stale = False
                try:
                    data_dt = datetime.strptime(as_of, "%Y-%m-%d")
                    diff_days = (now_dt - data_dt).days
                    if diff_days > 30:
                        is_stale = True
                except Exception:
                    pass

                if is_stale:
                    elig_status = "Stale - Verify"
                    elig_reason = "Attendance data is older than 30 days"
                elif pct >= threshold:
                    elig_status = "Eligible"
                    elig_reason = f"Attendance {pct:.1f}% meets required {threshold}% threshold"
                else:
                    elig_status = "Ineligible"
                    elig_reason = f"Attendance {pct:.1f}% is below required {threshold}% threshold"

            valid_records.append({
                "row": row_num,
                "rollNumber": roll,
                "studentName": name,
                "department": dept,
                "classesConducted": conducted_val,
                "classesAttended": attended_val,
                "attendancePercentage": pct,
                "dataAsOf": as_of,
                "eligibilityStatus": elig_status,
                "eligibilityReason": elig_reason
            })

    return {
        "totalRows": len(rows),
        "validRowsCount": len(valid_records),
        "invalidRowsCount": len(invalid_records),
        "unmatchedCount": len(unmatched_records),
        "duplicateCount": len(duplicates),
        "validRecords": valid_records,
        "invalidRecords": invalid_records,
        "threshold": threshold
    }

def commit_attendance_import(valid_records, filename, imported_by="In-Charge", notes=""):
    if not valid_records:
        raise ValueError("Cannot import: 0 valid records found.")

    batch_id = f"BATCH-{datetime.now().strftime('%Y%m%d%H%M%S')}-{uuid.uuid4().hex[:6].upper()}"
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    conn = get_db()
    cursor = conn.cursor()

    try:
        # Atomic commit
        cursor.execute("BEGIN TRANSACTION")

        for r in valid_records:
            roll = r["rollNumber"]
            conducted = r["classesConducted"]
            attended = r["classesAttended"]
            pct = r["attendancePercentage"]
            as_of = r["dataAsOf"]
            elig_status = r["eligibilityStatus"]
            is_stale = 1 if elig_status == "Stale - Verify" else 0
            is_missing = 1 if conducted == 0 or pct is None else 0

            cursor.execute("""
            INSERT OR REPLACE INTO attendance_records (
                roll_number, student_name, department, year_section,
                classes_conducted, classes_attended, attendance_percentage,
                data_as_of, missing_attendance, is_stale, eligibility_status,
                last_imported_at, import_batch_id
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                roll, r["studentName"], r["department"], "Current",
                conducted, attended, pct, as_of, is_missing, is_stale,
                elig_status, now_str, batch_id
            ))

        # Record import batch
        cursor.execute("""
        INSERT INTO attendance_imports (
            batch_id, filename, imported_by, total_rows, valid_rows,
            invalid_rows, unmatched_rows, duplicate_rows, imported_at, status, notes
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'SUCCESS', ?)
        """, (
            batch_id, filename, imported_by, len(valid_records),
            len(valid_records), 0, 0, 0, now_str, notes
        ))

        conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        raise RuntimeError(f"Database transaction failed; previously valid dataset preserved: {str(e)}")
    
    conn.close()
    return {
        "batchId": batch_id,
        "importedAt": now_str,
        "recordsCount": len(valid_records)
    }

def get_import_history():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
    SELECT batch_id, filename, imported_by, total_rows, valid_rows, invalid_rows, imported_at, status, notes
    FROM attendance_imports
    ORDER BY id DESC
    """)
    rows = [dict(r) for r in cursor.fetchall()]
    conn.close()
    return rows
